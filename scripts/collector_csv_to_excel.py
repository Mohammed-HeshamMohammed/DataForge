"""Turn a CSV downloaded by one of the browser collectors (Zillow, Amazon, real estate) into a clean Excel workbook.

The workbook opens on a Summary sheet (headline figures and breakdowns, as live formulas), followed by:

* Homes / Products / Listings - one row per item with the short fields only, as a filterable Excel table;
* Details - the long texts (description, facts, price history, ...) for the items whose page was collected;
* Photos - one row per photo link.

Long multi-line texts never sit in the main table, so its rows stay one line high. Numbers become numbers
(ID-like columns such as zpid, ASIN, zip and MLS # stay text), dates become dates, and links are clickable.

Example:
    python scripts/collector_csv_to_excel.py "%USERPROFILE%\\Downloads\\zillow_houston-tx_1235_homes_details.csv"
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT = "Arial"
EXCEL_CELL_LIMIT = 32_000  # Excel holds at most 32,767 characters in a cell
ID_COLUMNS = re.compile(r"zpid|asin|zip|postal|mls|ean|upc|isbn|listing id|reference|account|cause|doc id", re.IGNORECASE)
NUMBER = re.compile(r"-?\d+(\.\d+)?")
ISO_DATETIME = re.compile(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(:\d{2}(\.\d+)?)?Z?")
ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
ENUM_VALUE = re.compile(r"[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+|[A-Z]{3,}")  # SINGLE_FAMILY, TOWNHOUSE, LOT
UNKNOWN_IF_NEGATIVE = {"Days on Zillow"}  # Zillow writes -1 when it does not know
ENUM_COLUMNS = {"Home type", "Status", "Property type", "Purpose"}
HEADER_FILL = PatternFill("solid", start_color="1F3864")
SECTION_FILL = PatternFill("solid", start_color="D9E1F2")
LINK_FONT = Font(name=FONT, size=10, color="0563C1", underline="single")

# Number formats by column; money columns get the currency format chosen per file.
MONEY = {"Price", "Price (max)", "Zestimate", "Rent Zestimate", "Tax assessed value", "HOA fee (monthly)", "HOA fee", "List price",
         "Price per sqft", "Price per area", "Service charge", "Minimum bid", "Adjudged value", "Market value", "Appraised value",
         "Land value", "Building value", "Last sale price"}
FORMATS = {
    **dict.fromkeys(["Living area (sqft)", "Ratings count", "Days on Zillow", "Days on market", "Page views", "Favorites", "Area",
                     "Land area (sqft)"], "#,##0"),
    **dict.fromkeys(["Acreage"], "0.00"),
    **dict.fromkeys(["Year built"], "0"),
    **dict.fromkeys(["Lot area"], "#,##0.00"),
    **dict.fromkeys(["Latitude", "Longitude"], "0.000000"),
    **dict.fromkeys(["Rating"], "0.0"),
    **dict.fromkeys(["Property tax rate"], '0.00"%"'),
}
RENAMES = {"Detail URL": "Link", "Product URL": "Link", "Listing URL": "Link", "Image": "Photo"}


@dataclass
class Profile:
    site: str
    noun: str
    sheet: str
    title: str
    link: str
    order: list[str]
    long_text: list[str]
    photo_lists: list[str]
    drop: set[str] = field(default_factory=set)
    breakdowns: list[tuple[str, int]] = field(default_factory=list)
    stats: list[tuple[str, str, str]] = field(default_factory=list)
    money_format: str = "#,##0"
    link_text: str = ""  # text shown for the link column; default "View on <site>"
    extra_links: list[str] = field(default_factory=list)  # more link columns, shown as "Open"
    # columns computed in the workbook as live formulas: (name, formula with {Column} placeholders, number format)
    computed: list[tuple[str, str, str]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)  # extra lines under the Summary


PRICE_STATS = [("Lowest price", "MIN", "Price"), ("Median price", "MEDIAN", "Price"),
               ("Average price", "AVERAGE", "Price"), ("Highest price", "MAX", "Price")]
ZILLOW = Profile(
    site="Zillow", noun="homes", sheet="Homes", title="Address", link="Detail URL",
    order=["Address", "City", "State", "Zip", "Status", "Home type", "Price", "Zestimate", "Rent Zestimate", "Beds", "Baths",
           "Living area (sqft)", "Lot area", "Lot unit", "Year built", "Price per sqft", "HOA fee (monthly)", "Property tax rate",
           "Tax assessed value", "Days on Zillow", "Date posted", "Page views", "Favorites", "Brokerage", "MLS #", "MLS name",
           "Listing type", "Latitude", "Longitude", "Detail URL", "Image", "Page", "Position", "zpid"],
    long_text=["Description", "Facts & features", "Price history", "Tax history", "Schools"],
    photo_lists=["All photos"], drop={"Facts JSON", "Street"},
    breakdowns=[("Home type", 20), ("City", 15), ("Zip", 15)],
    stats=PRICE_STATS + [("Median Zestimate", "MEDIAN", "Zestimate"), ("Median living area (sqft)", "MEDIAN", "Living area (sqft)"),
                         ("Median days on Zillow", "MEDIAN", "Days on Zillow")],
    money_format="$#,##0",
)
AMAZON = Profile(
    site="Amazon", noun="products", sheet="Products", title="Title", link="Product URL",
    order=["Title", "Brand", "Price", "Currency", "List price", "Rating", "Ratings count", "Badge", "Sponsored", "Seller",
           "Fulfilled by", "Availability", "Delivery", "Category path", "Product URL", "Image", "Page", "Position", "ASIN"],
    long_text=["About this item", "Product details", "Features & Specs", "Additional details", "Product description",
               "From the manufacturer", "Best Sellers Rank"],
    photo_lists=["All images", "More images"], drop={"Details JSON"},
    breakdowns=[("Brand", 15), ("Seller", 15)],
    stats=PRICE_STATS + [("Average rating", "AVERAGE", "Rating"), ("Total ratings", "SUM", "Ratings count")],
    money_format="#,##0.00",
)
REAL_ESTATE = Profile(
    site="", noun="listings", sheet="Listings", title="Address", link="Listing URL",
    order=["Address", "City", "State", "Zip", "Status", "Property type", "Price", "Price (max)", "Beds", "Baths", "Living area (sqft)",
           "Lot size", "Year built", "HOA fee", "Days on market", "Listed", "Brokerage", "MLS #", "Latitude", "Longitude",
           "Listing URL", "Image", "Page", "Site", "Listing ID"],
    long_text=["Description", "Facts"],
    photo_lists=["Photos"], drop={"Facts JSON"},
    breakdowns=[("Property type", 20), ("City", 15), ("Zip", 15), ("Site", 10)],
    stats=PRICE_STATS + [("Median living area (sqft)", "MEDIAN", "Living area (sqft)"), ("Median beds", "MEDIAN", "Beds"),
                         ("Median days on market", "MEDIAN", "Days on market")],
    money_format="$#,##0",
)
TAX_SALES = Profile(
    site="Harris County Tax Office", noun="properties", sheet="Properties", title="Address", link="Sale details URL",
    order=["Sale date", "Address", "City", "State", "Zip", "Status", "Sale type", "Precinct", "Sale #", "Minimum bid",
           "Adjudged value", "Market value", "Property class", "Year built", "Living area (sqft)", "Land area (sqft)", "Bedrooms",
           "Full baths", "Half baths", "Quality", "Homestead", "Judgment date", "Tax years", "Cause #", "HCAD account",
           "HCAD link", "Sale details URL", "Image"],
    long_text=["Legal description"], photo_lists=[], link_text="Details", extra_links=["HCAD link"],
    breakdowns=[("Status", 6), ("Precinct", 10), ("Zip", 15)],
    stats=[("Lowest minimum bid", "MIN", "Minimum bid"), ("Median minimum bid", "MEDIAN", "Minimum bid"),
           ("Highest minimum bid", "MAX", "Minimum bid"), ("Median adjudged value", "MEDIAN", "Adjudged value"),
           ("Median bid / value", "MEDIAN", "Bid / value")],
    money_format="$#,##0",
    computed=[("Bid / value", '=IFERROR({Minimum bid}/{Adjudged value},"")', "0%"),
              ("Bid / market value", '=IFERROR({Minimum bid}/{Market value},"")', "0%")],
    notes=["Resales and EOS sales show no minimum bid or adjudged value on the Tax Office list, so those cells are blank.",
           "Source: Harris County Tax Office tax sale listing (hctax.net). Verify every property with the Tax Office before bidding."],
)
FORECLOSURE_NOTICES = Profile(
    site="Harris County Clerk", noun="notices", sheet="Notices", title="Doc ID", link="Notice URL",
    order=["Doc ID", "Sale date", "File date", "Pages", "Notice URL", "County"], long_text=[], photo_lists=[],
    link_text="Open notice", breakdowns=[("Sale date", 12)],
    notes=["Source: Harris County Clerk trustee foreclosure notices index. The property address is inside each notice."],
)
HCAD_ROLL = Profile(
    site="HCAD", noun="properties", sheet="Properties", title="Address", link="HCAD link",
    order=["HCAD account", "Address", "City", "Zip", "Property class", "Market area", "Neighborhood", "Year built",
           "Living area (sqft)", "Land area (sqft)", "Acreage", "Bedrooms", "Full baths", "Half baths", "Quality", "Market value",
           "Appraised value", "Land value", "Building value", "Prior market value", "Owner since", "Owner", "Owner type",
           "Mailing address", "Mailing city", "Mailing state", "Mailing zip", "Absentee owner", "Out of state owner",
           "Homestead", "Mail undeliverable", "Legal description", "HCAD link"],
    long_text=[], photo_lists=[], link_text="HCAD page",
    breakdowns=[("Property class", 10), ("Zip", 20), ("Owner type", 6), ("Absentee owner", 3), ("Homestead", 2),
                ("Out of state owner", 2)],
    stats=[("Median market value", "MEDIAN", "Market value"), ("Median year built", "MEDIAN", "Year built"),
           ("Median living area (sqft)", "MEDIAN", "Living area (sqft)"), ("Median years owned", "MEDIAN", "Years owned"),
           ("Median value change", "MEDIAN", "Value change")],
    money_format="$#,##0",
    computed=[("Years owned", '=IFERROR(DATEDIF({Owner since},TODAY(),"y"),"")', "0"),
              ("Value change", '=IFERROR({Market value}/{Prior market value}-1,"")', "0%")],
    notes=["Source: Harris Central Appraisal District public data files (hcad.org), 2026 certified values.",
           "Owner and mailing address are as HCAD publishes them. Absentee owner: the mailing address is not the property.",
           "Homestead: a residential homestead exemption is on the account. Other exemptions are left out on purpose.",
           "Years owned counts from the date HCAD records for the current owner."],
)
GENERIC = Profile(site="", noun="rows", sheet="Data", title="", link="", order=[], long_text=[], photo_lists=[])


def profile_for(header: list[str]) -> Profile:
    if "zpid" in header:
        return ZILLOW
    if "ASIN" in header:
        return AMAZON
    if "Listing URL" in header:
        return REAL_ESTATE
    if "Minimum bid" in header:
        return TAX_SALES
    if "Notice URL" in header:
        return FORECLOSURE_NOTICES
    if "Owner since" in header:
        return HCAD_ROLL
    return GENERIC


def cell_value(column: str, text: str):
    """A CSV text as the value Excel should hold: numbers, dates and Yes/No where they apply."""
    text = (text or "").strip()
    if text[:1] == "'" and text[1:2] in ("=", "+", "-", "@"):
        text = text[1:]  # undo the collector's guard against Excel formulas
    if not text:
        return None
    if ID_COLUMNS.search(column):
        return text[:EXCEL_CELL_LIMIT]
    lowered = text.lower()
    if lowered in ("true", "false"):
        return "Yes" if lowered == "true" else "No"
    if NUMBER.fullmatch(text):
        number = float(text) if "." in text else int(text)
        return None if column in UNKNOWN_IF_NEGATIVE and number < 0 else number
    if ISO_DATETIME.fullmatch(text):
        return datetime.fromisoformat(text.rstrip("Z")[:19])
    if ISO_DATE.fullmatch(text):
        return datetime.fromisoformat(text)
    if column in ENUM_COLUMNS and ENUM_VALUE.fullmatch(text):
        return text.replace("_", " ").capitalize()
    if lowered == "map" and column == "Page":
        return "Map"
    return text[:EXCEL_CELL_LIMIT]


def spread_specs(header: list[str], rows: list[dict]) -> list[str]:
    """Collector rows carry their specs as JSON (Amazon "Details JSON", real estate "Facts JSON"); spread the common
    ones into "Spec: <name>" / "Fact: <name>" columns."""
    source = next((name for name in ("Details JSON", "Facts JSON") if name in header), None)
    if not source:
        return []
    prefix = "Spec: " if source == "Details JSON" else "Fact: "
    parsed = []
    for row in rows:
        try:
            specs = json.loads(row.get(source) or "{}")
        except ValueError:
            specs = {}
        parsed.append(specs if isinstance(specs, dict) else {})
    counts = Counter(key for specs in parsed for key in specs)
    threshold = max(2, round(0.05 * sum(1 for specs in parsed if specs)))
    keys = [key for key, count in counts.most_common(80) if count >= threshold and prefix + key not in header]
    for row, specs in zip(rows, parsed):
        row.update({prefix + key: specs[key] for key in keys if key in specs})
    return [prefix + key for key in keys]


SITE_NAMES = {"realtor.com": "Realtor.com", "redfin.com": "Redfin", "trulia.com": "Trulia", "homes.com": "Homes.com",
              "apartments.com": "Apartments.com", "zillow.com": "Zillow", "har.com": "HAR.com",
              "coldwellbankerhomes.com": "Coldwell Banker", "remax.com": "RE/MAX", "weichert.com": "Weichert",
              "estately.com": "Estately", "opendoor.com": "Opendoor", "streeteasy.com": "StreetEasy",
              "windermere.com": "Windermere", "elliman.com": "Douglas Elliman", "hotpads.com": "HotPads",
              "apartmentguide.com": "ApartmentGuide", "apartmentfinder.com": "ApartmentFinder",
              "craigslist.org": "Craigslist", "landwatch.com": "LandWatch", "realtytrac.com": "RealtyTrac",
              "showcase.com": "Showcase", "newhomesource.com": "NewHomeSource", "century21.com": "Century 21"}


def site_name(host: str) -> str:
    host = host.lower().removeprefix("www.")
    domain = ".".join(host.split(".")[-2:])  # houston.craigslist.org -> craigslist.org
    return SITE_NAMES.get(host) or SITE_NAMES.get(domain) or (domain.split(".")[0].capitalize() if host else "")


def area_name(source: Path) -> str:
    """'zillow_houston-tx_1235_homes_details' -> 'Houston TX'."""
    parts = source.stem.split("_")
    if parts[0] == "zillow" and len(parts) >= 3 and parts[1] not in ("areas", "search"):
        return " ".join(word.upper() if len(word) == 2 else word.capitalize() for word in parts[1].split("-"))
    return ""


def style_header(sheet, columns: int, height: float = 30) -> None:
    for cell in sheet[1][:columns]:
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = height


def add_table(sheet, name: str, columns: int, rows: int) -> None:
    table = Table(displayName=name, ref=f"A1:{get_column_letter(columns)}{rows + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    sheet.add_table(table)


def fit_widths(sheet, header: list[str], values: list[list], fixed: dict[str, float] | None = None, cap: float = 40) -> None:
    for index, name in enumerate(header, 1):
        if fixed and name in fixed:
            width = fixed[name]
        else:
            sample = [values[r][index - 1] for r in range(min(len(values), 300))]
            longest = max([len(str(v).split("\n")[0]) for v in sample if v is not None] + [0])
            width = min(cap, max(len(name) * 0.9 + 4, longest * 1.05 + 2, 9))
        sheet.column_dimensions[get_column_letter(index)].width = width


def build(source: Path, out: Path) -> dict:
    with source.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        header = list(reader.fieldnames or [])
        raw = [row for row in reader if any((value or "").strip() for value in row.values())]
    if not header:
        raise SystemExit(f"{source} is empty")
    profile = profile_for(header)
    header += spread_specs(header, raw)
    rows = [{column: cell_value(column, row.get(column) or "") for column in header} for row in raw]

    site = profile.site or site_name(next((row.get("Site") for row in rows if row.get("Site")), "")) or "Collected"
    money = profile.money_format
    currencies = {row.get("Currency") for row in rows if row.get("Currency")}
    single_currency = len(currencies) == 1 and next(iter(currencies))
    drop = set(profile.drop) | ({"Currency"} if single_currency else set())
    if single_currency:
        money = f'#,##0{".00" if profile is AMAZON else ""} "{single_currency}"'

    title_column = profile.title if profile.title in header else header[0]
    link_column = profile.link if profile.link in header else None
    side = set(profile.long_text) | set(profile.photo_lists) | drop
    ordered = [c for c in profile.order if c in header] + [c for c in header if c not in profile.order]
    main_columns = [c for c in ordered if c not in side and any(row.get(c) not in (None, "") for row in rows)]

    computed = [(name, formula, number_format) for name, formula, number_format in profile.computed
                if all(part in main_columns for part in re.findall(r"\{([^}]+)\}", formula))]
    link_label = profile.link_text or f"View on {site}"

    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    summary = workbook.active
    summary.title = "Summary"

    # Main table: short fields only, one line per row.
    data = workbook.create_sheet(profile.sheet)
    shown = [RENAMES.get(c, c) for c in main_columns] + [name for name, _, _ in computed]
    data.append(shown)
    values = []
    for row in rows:
        line = []
        for column in main_columns:
            value = row.get(column)
            if column == link_column and value:
                value = link_label
            elif column in profile.extra_links and value:
                value = "Open"
            elif column == "Image" and value:
                value = "Photo"
            line.append(value)
        values.append(line)
        data.append(line)
    for row_index, row in enumerate(rows, 2):
        for column_index, column in enumerate(main_columns, 1):
            cell = data.cell(row_index, column_index)
            cell.font = Font(name=FONT, size=10)
            if column in (link_column, "Image", *profile.extra_links) and row.get(column):
                cell.hyperlink = row[column]
                cell.font = LINK_FONT
            elif column in MONEY:
                cell.number_format = money
            elif column in FORMATS:
                cell.number_format = FORMATS[column]
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm" if column.endswith(" at") else "yyyy-mm-dd"
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # collected text must stay text, never become a formula
        for offset, (name, formula, number_format) in enumerate(computed):
            cell = data.cell(row_index, len(main_columns) + 1 + offset)
            cell.value = re.sub(r"\{([^}]+)\}", lambda m: f"{get_column_letter(main_columns.index(m.group(1)) + 1)}{row_index}", formula)
            cell.number_format = number_format
            cell.font = Font(name=FONT, size=10)
    values = [line + [None] * len(computed) for line in values]
    style_header(data, len(shown))
    add_table(data, profile.sheet, len(shown), len(rows))
    fit_widths(data, shown, values, {"Link": 16, "Photo": 9, **{name: 12 for name, _, _ in computed}})
    data.freeze_panes = "B2"

    # Long texts, for the items that have them.
    long_columns = [c for c in profile.long_text if c in header and any(row.get(c) for row in rows)]
    detail_rows = [row for row in rows if any(row.get(c) for c in long_columns)]
    if detail_rows:
        details = workbook.create_sheet("Details")
        detail_header = [title_column] + (["Link"] if link_column else []) + long_columns
        details.append(detail_header)
        for index, row in enumerate(detail_rows, 2):
            details.append([row.get(title_column)] + ([link_label] if link_column else []) + [row.get(c) for c in long_columns])
            for cell in details[index]:
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if link_column and row.get(link_column):
                details.cell(index, 2).hyperlink = row[link_column]
                details.cell(index, 2).font = LINK_FONT
        style_header(details, len(detail_header))
        add_table(details, "Details", len(detail_header), len(detail_rows))
        for index, name in enumerate(detail_header, 1):
            details.column_dimensions[get_column_letter(index)].width = 38 if index == 1 else 16 if name == "Link" else 70
        details.freeze_panes = "B2"

    # One row per photo.
    photo_column = next((c for c in profile.photo_lists if c in header and any(row.get(c) for row in rows)), None)
    photo_rows = []
    if photo_column:
        for row in rows:
            for number, url in enumerate(re.split(r"\s+", str(row.get(photo_column) or "").strip()), 1):
                if url.startswith("http"):
                    photo_rows.append([row.get(title_column), number, url])
    if photo_rows:
        photos = workbook.create_sheet("Photos")
        photos.append([title_column, "Photo #", "Photo link"])
        for index, line in enumerate(photo_rows, 2):
            photos.append(line)
            for cell in photos[index]:
                cell.font = Font(name=FONT, size=10)
            photos.cell(index, 3).hyperlink = line[2]
            photos.cell(index, 3).font = LINK_FONT
        style_header(photos, 3)
        add_table(photos, "Photos", 3, len(photo_rows))
        for letter, width in zip("ABC", (45, 9, 100)):
            photos.column_dimensions[letter].width = width
        photos.freeze_panes = "A2"

    write_summary(summary, profile, site, source, main_columns + [name for name, _, _ in computed], rows, money,
                  {name: number_format for name, _, number_format in computed})
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return {"rows": len(rows), "columns": len(main_columns), "details": len(detail_rows), "photos": len(photo_rows),
            "sheets": workbook.sheetnames}


def write_summary(sheet, profile: Profile, site: str, source: Path, columns: list[str], rows: list[dict], money: str,
                  formats: dict[str, str] | None = None) -> None:
    last = len(rows) + 1
    data_name = profile.sheet

    def column_range(name: str) -> str | None:
        if name not in columns:
            return None
        letter = get_column_letter(columns.index(name) + 1)
        return f"{data_name}!${letter}$2:${letter}${last}"

    def label(cell, text, bold=False, size=10, color=None, italic=False):
        cell.value = text
        cell.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

    noun = profile.noun if profile is not GENERIC else "rows"
    area = area_name(source)
    label(sheet["A1"], f"{site} {noun}" + (f" — {area}" if area else ""), bold=True, size=16, color="1F3864")
    label(sheet["A2"], f"Made from {source.name} on {datetime.now():%Y-%m-%d %H:%M}", italic=True, color="7F7F7F")

    row = 4
    label(sheet.cell(row, 1), "Key figures", bold=True, size=12, color="1F3864")
    row += 1
    count_range = f"{data_name}!$A$2:$A${last}"
    figures = [(f"{noun.capitalize()} collected", f"=COUNTA({count_range})", "#,##0")]
    page_range = column_range("Page")
    if page_range:
        figures.append(("From result pages", f"=COUNT({page_range})", "#,##0"))
        if any(r.get("Page") == "Map" for r in rows):
            figures.append(("From the map", f'=COUNTIF({page_range},"Map")', "#,##0"))
    detail_range = column_range("Details collected at")
    if detail_range:
        figures.append(("With full details", f"=COUNTA({detail_range})", "#,##0"))
    for text, function, column in profile.stats:
        target = column_range(column)
        if target:
            number_format = money if column in MONEY else (formats or {}).get(column) or FORMATS.get(column, "#,##0")
            figures.append((text, f'=IFERROR({function}({target}),"–")', number_format))
    total_cell = f"$B${row}"
    for text, formula, number_format in figures:
        label(sheet.cell(row, 1), text)
        cell = sheet.cell(row, 2)
        cell.value = formula
        cell.number_format = number_format
        cell.font = Font(name=FONT, size=10, bold=text.endswith("collected"))
        cell.alignment = Alignment(horizontal="right")
        row += 1

    price_range = column_range("Price")
    rating_range = column_range("Rating")
    for column, limit in profile.breakdowns:
        target = column_range(column)
        if not target:
            continue
        common = Counter(r[column] for r in rows if r.get(column) not in (None, "")).most_common(limit)
        if len(common) < 2:
            continue
        row += 1
        heads = [RENAMES.get(column, column), noun.capitalize(), "Share"] + (["Average price"] if price_range else []) + \
                (["Average rating"] if rating_range else [])
        label(sheet.cell(row, 1), f"By {column.lower()}" + (f" (top {limit})" if len(common) == limit else ""),
              bold=True, size=12, color="1F3864")
        row += 1
        for offset, text in enumerate(heads):
            cell = sheet.cell(row, 1 + offset)
            label(cell, text, bold=True)
            cell.fill = SECTION_FILL
            cell.alignment = Alignment(horizontal="left" if offset == 0 else "right")
        row += 1
        for value, _ in common:
            criteria = '"=' + re.sub(r"([~*?])", r"~\1", str(value)).replace('"', '""') + '"'
            label(sheet.cell(row, 1), value)
            cells = [(f"=COUNTIF({target},{criteria})", "#,##0"), (f"=IFERROR(B{row}/{total_cell},0)", "0.0%")]
            if price_range:
                cells.append((f'=IFERROR(AVERAGEIF({target},{criteria},{price_range}),"–")', money))
            if rating_range:
                cells.append((f'=IFERROR(AVERAGEIF({target},{criteria},{rating_range}),"–")', "0.0"))
            for offset, (formula, number_format) in enumerate(cells, 2):
                cell = sheet.cell(row, offset)
                cell.value = formula
                cell.number_format = number_format
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(horizontal="right")
            row += 1

    row += 1
    notes = [f"All figures are live formulas over the {data_name} sheet.",
             "Each breakdown lists the most common values when this file was made."]
    if any(r.get("Page") == "Map" for r in rows):
        notes.append('"Map" in the Page column marks homes collected from the map pins (they have no list position).')
    notes += profile.notes
    for note in notes:
        label(sheet.cell(row, 1), note, italic=True, color="7F7F7F")
        row += 1
    for letter, width in zip("ABCDE", (34, 16, 10, 18, 15)):
        sheet.column_dimensions[letter].width = width
    sheet.sheet_view.showGridLines = False


def main(argv: list[str] | None = None) -> Path:
    parser = argparse.ArgumentParser(description="Convert a collector CSV to a formatted Excel workbook.")
    parser.add_argument("csv", type=Path, help="the CSV the collector bookmark downloaded")
    parser.add_argument("--out", type=Path, help="output .xlsx (default: next to the CSV)")
    args = parser.parse_args(argv)
    if not args.csv.is_file():
        parser.error(f"no such file: {args.csv}")
    out = args.out or args.csv.with_suffix(".xlsx")
    result = build(args.csv, out)
    print(f"wrote {result['rows']} rows ({', '.join(result['sheets'])}) to {out}", file=sys.stderr)
    return out


if __name__ == "__main__":
    main()
