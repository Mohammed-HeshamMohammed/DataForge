"""Export Amazon search results to Excel, through Amazon's official Creators API or a collector CSV.

Two ways in:

* An API search: give an Amazon search link (or keywords / a category browse node). One API search returns
  at most 100 products; --price-slices splits it into price ranges of up to 100 products each, so a large
  category (1,000+ products) can be collected in full.
* Products collected in the browser: give the CSV that the bookmark from scripts/amazon_search_collector.html
  downloads (--from-collector). Its search data (price, rating, ratings count, badges) and, when collected,
  each product page's details (Product details, About this item, Features & Specs, description) are exported
  as is, with the specs most products share spread into "Spec: <name>" columns. When API credentials are
  set, every product is also enriched with brand, features, images, seller and sales rank through GetItems.

The script never loads amazon.* web pages: product data comes from the Creators API or from that CSV.

The Creators API replaced the Product Advertising API 5.0 (retired 15 May 2026). It needs an Amazon
Associates account eligible for the API (10 qualifying sales in the last 30 days) and credentials created in
Associates Central, passed as environment variables:

    CREATORS_CREDENTIAL_ID      Credential ID
    CREATORS_CREDENTIAL_SECRET  Credential Secret
    CREATORS_PARTNER_TAG        your store/tracking ID for that marketplace, e.g. mystore-21
    CREATORS_TOKEN_URL          optional: token endpoint, if Associates Central shows a different one

New accounts may make 1 request per second and 8,640 per day. The API has no customer ratings or review
counts; those come from the collector CSV only.

Examples:
    python scripts/amazon_creators_export.py "https://www.amazon.eg/s?i=fashion&rh=n%3A21833439031&s=review-rank" --price-slices
    python scripts/amazon_creators_export.py --keywords "running shoes" --marketplace www.amazon.eg --variations
    python scripts/amazon_creators_export.py --from-collector "%USERPROFILE%\\Downloads\\amazon_eg_1008_products.csv"
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

API_BASE = "https://creatorsapi.amazon/catalog/v1"
CREDENTIAL_VARIABLES = ("CREATORS_CREDENTIAL_ID", "CREATORS_CREDENTIAL_SECRET", "CREATORS_PARTNER_TAG")
# Login with Amazon token endpoints per credential region (Creators API docs, "Using cURL").
TOKEN_URLS = {
    "NA": "https://api.amazon.com/auth/o2/token",
    "EU": "https://api.amazon.co.uk/auth/o2/token",
    "FE": "https://api.amazon.co.jp/auth/o2/token",
}
# Every other marketplace (UK, EU, Egypt, India, UAE, Saudi Arabia, Turkey, ...) is in the EU region.
REGION_MARKETPLACES = {
    "NA": {"www.amazon.com", "www.amazon.ca", "www.amazon.com.mx", "www.amazon.com.br"},
    "FE": {"www.amazon.co.jp", "www.amazon.sg", "www.amazon.com.au"},
}
# Amazon search-page `s=` values and the API sortBy they correspond to.
SORT_BY_URL_VALUE = {
    "review-rank": "AvgCustomerReviews",
    "price-asc-rank": "Price:LowToHigh",
    "price-desc-rank": "Price:HighToLow",
    "date-desc-rank": "NewestArrivals",
    "featured-rank": "Featured",
    "relevancerank": "Relevance",
    "relevanceblender": "Relevance",
}
SORT_CHOICES = sorted(set(SORT_BY_URL_VALUE.values()))
ITEM_RESOURCES = [
    "browseNodeInfo.browseNodes",
    "browseNodeInfo.websiteSalesRank",
    "images.primary.large",
    "images.variants.large",
    "itemInfo.byLineInfo",
    "itemInfo.classifications",
    "itemInfo.externalIds",
    "itemInfo.features",
    "itemInfo.manufactureInfo",
    "itemInfo.productInfo",
    "itemInfo.technicalInfo",
    "itemInfo.title",
    "offersV2.listings.availability",
    "offersV2.listings.condition",
    "offersV2.listings.dealDetails",
    "offersV2.listings.isBuyBoxWinner",
    "offersV2.listings.merchantInfo",
    "offersV2.listings.price",
    "parentASIN",
]
MAX_SEARCH_PAGES = 10  # the API serves itemPage 1-10 of 10 items only
MAX_RESULTS_PER_SEARCH = 100
MAX_VARIATION_PAGES = 10
CURRENCY_MINOR_UNITS = {"JPY": 0}  # minPrice/maxPrice are in the lowest denomination; 2 decimals otherwise
# Hints for the error reasons a first run most often hits.
ERROR_HINTS = {
    "AssociateNotEligible": "the Associates account needs 10 qualifying sales in the last 30 days to use the API",
    "InvalidPartnerTag": "CREATORS_PARTNER_TAG is not a tracking ID of this marketplace's Associates store",
    "InvalidAssociate": "the credentials are not linked to CREATORS_PARTNER_TAG",
    "InvalidClient": "CREATORS_CREDENTIAL_ID / CREATORS_CREDENTIAL_SECRET are wrong",
    "UnsupportedClient": "these credentials are for another region; set CREATORS_TOKEN_URL to the endpoint Associates Central shows",
}
# Workbook columns, in order; columns no product has a value for are left out of the sheet. The collector's
# per-product "Spec: <name>" columns go in front of "Page".
COLUMNS = [
    "ASIN", "Parent ASIN", "Title", "Brand", "Manufacturer", "Variation", "Color", "Size",
    "Price", "Currency", "Price (display)", "List price", "Saving %", "Rating", "Ratings count", "Badge",
    "Sponsored", "Availability", "Condition", "Seller", "Fulfilled by", "Deal", "Delivery",
    "Features", "Product group", "Binding", "Model", "EAN", "Category", "Category path", "Category sales rank",
    "Site sales rank", "Best Sellers Rank", "About this item", "Product details", "Features & Specs",
    "Additional details", "Product description", "From the manufacturer", "Details note", "Details collected at",
    "Page", "Position", "Image", "More images", "All images", "Product URL",
]
# Collector CSV columns that hold numbers.
COLLECTOR_NUMBERS = {"Price", "List price", "Rating", "Ratings count", "Page", "Position"}
WIDE_TEXT_COLUMNS = {
    "Title": 60, "Features": 80, "More images": 60, "Product URL": 45, "Image": 45, "Delivery": 40, "Category path": 50,
    "Best Sellers Rank": 50, "About this item": 80, "Product details": 50, "Features & Specs": 60,
    "Additional details": 50, "Product description": 80, "From the manufacturer": 60, "All images": 50,
}
WRAPPED_COLUMNS = {"Features", "More images"}
MAX_SPEC_COLUMNS = 80
EXCEL_CELL_LIMIT = 32_000  # Excel holds at most 32,767 characters in a cell


class ApiError(RuntimeError):
    def __init__(self, status: int, message: str) -> None:
        super().__init__(message)
        self.status = status


def lower_keys(value):
    """Lower-case every JSON key, so lookups match whether the API answers in camelCase or PascalCase."""
    if isinstance(value, dict):
        return {key.lower(): lower_keys(item) for key, item in value.items()}
    if isinstance(value, list):
        return [lower_keys(item) for item in value]
    return value


def dig(value, path: str):
    """Follow a dotted camelCase path through a lower_keys() result; None when any step is missing."""
    for part in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part.lower())
    return value


def describe_error(status: int, text: str) -> str:
    hints = [hint for reason, hint in ERROR_HINTS.items() if reason in text]
    return f"HTTP {status}: {text[:500]}" + "".join(f"\n  -> {hint}" for hint in hints)


def credentials() -> tuple[str, str, str] | None:
    values = tuple(os.environ.get(name, "") for name in CREDENTIAL_VARIABLES)
    return values if all(values) else None


class CreatorsApi:
    def __init__(self, credential_id: str, credential_secret: str, partner_tag: str, marketplace: str,
                 http: httpx.Client, token_url: str | None = None, min_interval: float = 1.1,
                 raw_log=None) -> None:
        region = next((name for name, hosts in REGION_MARKETPLACES.items() if marketplace in hosts), "EU")
        self.credential_id, self.credential_secret = credential_id, credential_secret
        self.partner_tag, self.marketplace, self.http = partner_tag, marketplace, http
        self.token_url = token_url or TOKEN_URLS[region]
        self.min_interval, self.raw_log = min_interval, raw_log
        self._token: str | None = None
        self._token_expiry = 0.0
        self._last_call = 0.0

    def _access_token(self) -> str:
        if self._token and time.monotonic() < self._token_expiry - 60:
            return self._token
        form = {
            "grant_type": "client_credentials",
            "client_id": self.credential_id,
            "client_secret": self.credential_secret,
            "scope": "creatorsapi::default",
        }
        response = self.http.post(self.token_url, data=form)
        if response.status_code in (400, 415):  # the docs show both a form and a JSON body; accept either
            response = self.http.post(self.token_url, json=form)
        if response.status_code != 200:
            raise ApiError(response.status_code, "token request failed: " + describe_error(response.status_code, response.text))
        body = response.json()
        self._token = body["access_token"]
        self._token_expiry = time.monotonic() + int(body.get("expires_in", 3600))
        return self._token

    def call(self, operation: str, payload: dict) -> dict:
        body = {"partnerTag": self.partner_tag, "marketplace": self.marketplace, **payload}
        for attempt in range(6):
            wait = self._last_call + self.min_interval - time.monotonic()
            if wait > 0:
                time.sleep(wait)
            headers = {"Authorization": f"Bearer {self._access_token()}", "x-marketplace": self.marketplace}
            self._last_call = time.monotonic()
            response = self.http.post(f"{API_BASE}/{operation}", json=body, headers=headers)
            if response.status_code == 401 and attempt == 0:
                self._token = None  # expired or revoked token: fetch a new one once
                continue
            if response.status_code == 429 or response.status_code >= 500:
                time.sleep(min(2 ** attempt * self.min_interval, 30))
                continue
            break
        if self.raw_log:
            self.raw_log.write(json.dumps({"operation": operation, "request": payload, "status": response.status_code,
                                           "response": response.text}, ensure_ascii=False) + "\n")
        if response.status_code != 200:
            raise ApiError(response.status_code, f"{operation} failed: " + describe_error(response.status_code, response.text))
        return lower_keys(response.json())


def query_from_url(url: str) -> tuple[str, dict, list[str]]:
    """Turn an Amazon search link into (marketplace, SearchItems parameters, notes about what was dropped)."""
    parsed = urlparse(url)
    host = (parsed.hostname or "").lower()
    if "amazon." not in host:
        raise SystemExit(f"not an Amazon link: {url}")
    marketplace = host if host.startswith("www.") else f"www.{host}"
    params = parse_qs(parsed.query)
    query: dict = {}
    notes: list[str] = []
    if params.get("k"):
        query["keywords"] = params["k"][0]
    refinements = params.get("rh", [""])[0]
    nodes = re.findall(r"(?:^|,)n:(\d+)", refinements)
    if nodes:
        query["browseNodeId"] = nodes[-1]
        if len(nodes) > 1:
            notes.append(f"the link filters by {len(nodes)} categories; the API takes one, using the last ({nodes[-1]})")
    elif params.get("node"):
        query["browseNodeId"] = params["node"][0]
    price = re.search(r"p_36:(\d*)-(\d*)", refinements)
    if price and price.group(1):
        query["minPrice"] = int(price.group(1))
    if price and price.group(2):
        query["maxPrice"] = int(price.group(2))
    sort = params.get("s", [""])[0]
    if sort in SORT_BY_URL_VALUE:
        query["sortBy"] = SORT_BY_URL_VALUE[sort]
    elif sort:
        notes.append(f"sort '{sort}' has no API equivalent; using the API's default order")
    if params.get("language", [""])[0].startswith("ar"):
        query["languagesOfPreference"] = ["ar_AE"]
    other = sorted({part.split(":")[0] for part in refinements.split(",") if part and not part.startswith(("n:", "p_36:"))})
    if other:
        notes.append(f"filters {', '.join(other)} have no API equivalent and were ignored")
    return marketplace, query, notes


def search_page(api: CreatorsApi, query: dict, page: int) -> dict | None:
    """One page of search results, or None when there are no results on it."""
    try:
        data = api.call("searchItems", {**query, "itemCount": 10, "itemPage": page, "resources": ITEM_RESOURCES})
    except ApiError as error:
        if error.status == 404:
            return None
        raise
    return data.get("searchresult") or None


def search(api: CreatorsApi, query: dict, pages: int, first: dict | None = None) -> tuple[list[dict], int | None]:
    """Every result page of one search (at most 100 products); `first` is page 1 when already fetched."""
    items: list[dict] = []
    total = None
    for page in range(1, pages + 1):
        result = first if page == 1 and first else search_page(api, query, page)
        if not result:
            break
        total = result.get("totalresultcount", total)
        batch = result.get("items") or []
        items.extend(batch)
        if len(batch) < 10 or (total is not None and len(items) >= total):
            break
    return items, total


def buy_box(item: dict) -> dict:
    listings = dig(item, "offersV2.listings") or []
    return next((listing for listing in listings if listing.get("isbuyboxwinner")), listings[0] if listings else {})


def search_by_price_slices(api: CreatorsApi, query: dict, pages: int) -> tuple[list[dict], int | None]:
    """Collect a search with more results than one search returns by splitting it into price ranges.

    Each range with more than 100 results is halved until every range fits in one search. Products without a
    price are only found by the first, unsliced search.
    """
    items, total = search(api, query, pages)
    print(f"  without price ranges: {len(items)} of {total} products", file=sys.stderr)
    if total is None or total <= len(items):
        return items, total
    found = {item["asin"]: item for item in items}
    top = (search_page(api, {**query, "sortBy": "Price:HighToLow"}, 1) or {}).get("items") or []
    offers = [buy_box(item) for item in top]
    prices = [amount for offer in offers if (amount := dig(offer, "price.money.amount"))]
    currency = next((code for offer in offers if (code := dig(offer, "price.money.currency"))), None)
    unit = 10 ** CURRENCY_MINOR_UNITS.get(currency, 2)
    low = query.get("minPrice", 1)
    high = query.get("maxPrice") or (int(max(prices) * unit) + unit if prices else 10 ** 9)
    ranges = [(low, high)]
    while ranges:
        low, high = ranges.pop(0)
        sliced = {**query, "minPrice": low, "maxPrice": high}
        first = search_page(api, sliced, 1)
        if not first:
            continue
        count = first.get("totalresultcount") or 0
        if count > MAX_RESULTS_PER_SEARCH and high - low > unit:
            middle = (low + high) // 2
            ranges[:0] = [(low, middle), (middle + 1, high)]
            continue
        batch, _ = search(api, sliced, pages, first)
        new = [item for item in batch if item["asin"] not in found]
        found.update((item["asin"], item) for item in new)
        print(f"  price {low / unit:g}-{high / unit:g}: {len(batch)} of {count}, {len(new)} new, {len(found)} so far",
              file=sys.stderr)
    return list(found.values()), total


def expand_variations(api: CreatorsApi, items: list[dict]) -> list[dict]:
    """Replace each product that has variations with every variation of its family (each size/colour)."""
    expanded: list[dict] = []
    done_parents: set[str] = set()
    for item in items:
        parent = item.get("parentasin")
        if not parent:
            expanded.append(item)
            continue
        if parent in done_parents:
            continue
        done_parents.add(parent)
        family: list[dict] = []
        page = page_count = 1
        try:
            while page <= min(page_count, MAX_VARIATION_PAGES):
                data = api.call("getVariations", {"asin": item["asin"], "variationCount": 10, "variationPage": page,
                                                  "resources": ITEM_RESOURCES + ["variationSummary.variationDimension"]})
                result = data.get("variationsresult") or {}
                page_count = dig(result, "variationSummary.pageCount") or 1
                family.extend(result.get("items") or [])
                page += 1
        except ApiError as error:
            print(f"  variations of {item['asin']} skipped: {error}", file=sys.stderr)
        expanded.extend(family or [item])
        print(f"  {item['asin']}: {len(family)} variations", file=sys.stderr)
    return expanded


def get_items(api: CreatorsApi, asins: list[str]) -> dict[str, dict]:
    """Product details by ASIN, 10 per request; ASINs the API cannot return are left out."""
    details: dict[str, dict] = {}
    for start in range(0, len(asins), 10):
        batch = asins[start:start + 10]
        try:
            data = api.call("getItems", {"itemIds": batch, "itemIdType": "ASIN", "resources": ITEM_RESOURCES})
        except ApiError as error:
            if error.status != 404:
                raise
            data = {}
        result = data.get("itemsresult") or data.get("itemresults") or {}
        details.update((item["asin"], item) for item in result.get("items") or [])
        print(f"  details: {start + len(batch)}/{len(asins)} ASINs, {len(details)} found", file=sys.stderr)
    return details


def read_collector_csv(path: Path) -> list[dict]:
    """Rows of a CSV downloaded by the amazon_search_collector bookmark, with numbers as numbers."""
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    if not rows or "ASIN" not in rows[0]:
        raise SystemExit(f"{path} is not a collector CSV (it has no ASIN column)")
    for row in rows:
        for key, value in list(row.items()):
            value = (value or "").strip()
            if value[:1] == "'" and value[1:2] in ("=", "+", "-", "@"):
                value = value[1:]  # undo the collector's guard against Excel formulas
            if key in COLLECTOR_NUMBERS and value:
                try:
                    value = float(value) if "." in value else int(value)
                except ValueError:
                    pass
            row[key] = value or None
    unique = {row["ASIN"]: row for row in reversed(rows) if row.get("ASIN")}
    return list(reversed(unique.values()))


def joined(values, separator: str = "\n") -> str | None:
    values = [str(value) for value in values or [] if value not in (None, "")]
    return separator.join(values) or None


def product_row(item: dict) -> dict:
    offer = buy_box(item)
    nodes = dig(item, "browseNodeInfo.browseNodes") or [{}]
    return {
        "ASIN": item.get("asin"),
        "Parent ASIN": item.get("parentasin"),
        "Title": dig(item, "itemInfo.title.displayValue"),
        "Brand": dig(item, "itemInfo.byLineInfo.brand.displayValue"),
        "Manufacturer": dig(item, "itemInfo.byLineInfo.manufacturer.displayValue"),
        "Variation": joined((f"{a.get('name')}: {a.get('value')}" for a in item.get("variationattributes") or []), ", "),
        "Color": dig(item, "itemInfo.productInfo.color.displayValue"),
        "Size": dig(item, "itemInfo.productInfo.size.displayValue"),
        "Price": dig(offer, "price.money.amount"),
        "Currency": dig(offer, "price.money.currency"),
        "Price (display)": dig(offer, "price.money.displayAmount"),
        "Saving %": dig(offer, "price.savings.percentage"),
        "Availability": dig(offer, "availability.type"),
        "Condition": dig(offer, "condition.value"),
        "Seller": dig(offer, "merchantInfo.name"),
        "Deal": dig(offer, "dealDetails.badge"),
        "Features": joined(dig(item, "itemInfo.features.displayValues")),
        "Product group": dig(item, "itemInfo.classifications.productGroup.displayValue"),
        "Binding": dig(item, "itemInfo.classifications.binding.displayValue"),
        "Model": dig(item, "itemInfo.manufactureInfo.model.displayValue"),
        "EAN": joined(dig(item, "itemInfo.externalIds.eans.displayValues"), ", "),
        "Category": nodes[0].get("displayname"),
        "Category sales rank": nodes[0].get("salesrank"),
        "Site sales rank": dig(item, "browseNodeInfo.websiteSalesRank.salesRank"),
        "Image": dig(item, "images.primary.large.url"),
        "More images": joined(dig(variant, "large.url") for variant in dig(item, "images.variants") or []),
        "Product URL": item.get("detailpageurl"),
    }


def merge_collected(row: dict, item: dict | None) -> dict:
    """A collector row with the API's details on top; rating, badges and page position stay from the page."""
    if not item:
        return row
    return {**row, **{key: value for key, value in product_row(item).items() if value not in (None, "")}}


def add_spec_columns(rows: list[dict]) -> list[str]:
    """Spread each row's collector "Details JSON" into "Spec: <name>" values, for specs enough products share.

    A spec becomes a column when at least 5% of the products with details (and at least 2) have it, so
    one-off specs stay only in the "Product details" / "Features & Specs" text.
    """
    parsed = []
    for row in rows:
        try:
            specs = json.loads(row.pop("Details JSON", None) or "{}")
        except ValueError:
            specs = {}
        parsed.append(specs if isinstance(specs, dict) else {})
    counts = Counter(key for specs in parsed for key in specs)
    threshold = max(2, round(0.05 * sum(1 for specs in parsed if specs)))
    keys = [key for key, count in counts.most_common(MAX_SPEC_COLUMNS) if count >= threshold]
    for row, specs in zip(rows, parsed):
        row.update({f"Spec: {key}": specs[key] for key in keys if key in specs})
    return [f"Spec: {key}" for key in keys]


def write_workbook(path: Path, rows: list[dict], run_info: dict, spec_columns: list[str] | None = None) -> None:
    order = COLUMNS[:COLUMNS.index("Page")] + (spec_columns or []) + COLUMNS[COLUMNS.index("Page"):]
    columns = [column for column in order if any(row.get(column) not in (None, "") for row in rows)] or COLUMNS
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(columns)
    for row in rows:
        sheet.append([value[:EXCEL_CELL_LIMIT] if isinstance(value := row.get(column), str) else value for column in columns])
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"  # product text must stay text, never become a formula
            header = columns[cell.column - 1]
            if header in ("Product URL", "Image") and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
            if header in WRAPPED_COLUMNS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = WIDE_TEXT_COLUMNS.get(cell.value, max(12, len(cell.value) + 2))
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions

    info = workbook.create_sheet("Run info")
    for key, value in run_info.items():
        info.append([key, value if isinstance(value, (int, float)) or value is None else str(value)])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Amazon search results to Excel via the Creators API.")
    parser.add_argument("url", nargs="?", help="an Amazon search link, e.g. https://www.amazon.eg/s?k=...")
    parser.add_argument("--keywords", help="search words (override the link's)")
    parser.add_argument("--browse-node", help="category browse node ID (overrides the link's)")
    parser.add_argument("--marketplace", help="e.g. www.amazon.eg (default: the link's site, else www.amazon.eg)")
    parser.add_argument("--sort", choices=SORT_CHOICES, help="result order (overrides the link's)")
    parser.add_argument("--search-index", help="API category index, e.g. Electronics (Egypt has no Fashion index)")
    parser.add_argument("--pages", type=int, default=MAX_SEARCH_PAGES, help="result pages of 10 products, 1-10 (default 10)")
    parser.add_argument("--price-slices", action="store_true", help="split the search by price to get past 100 products")
    parser.add_argument("--variations", action="store_true", help="also export every size/colour of each product")
    parser.add_argument("--from-collector", type=Path, metavar="CSV", help="export a CSV downloaded by the collector bookmark")
    parser.add_argument("--no-api", action="store_true", help="with --from-collector: export the page data without API details")
    parser.add_argument("--out", type=Path, help="output .xlsx (default: Data/amazon/<site>_<time>.xlsx)")
    args = parser.parse_args(argv)
    if args.from_collector:
        if args.url or args.keywords or args.browse_node or args.price_slices or args.variations:
            parser.error("--from-collector exports the CSV's products; it does not take a search or --price-slices/--variations")
        if not args.from_collector.is_file():
            parser.error(f"no such file: {args.from_collector}")
    elif not (args.url or args.keywords or args.browse_node):
        parser.error("give an Amazon search link, --keywords, --browse-node, or --from-collector")
    if not 1 <= args.pages <= MAX_SEARCH_PAGES:
        parser.error(f"--pages must be between 1 and {MAX_SEARCH_PAGES}")
    return args


def main(argv: list[str] | None = None, http: httpx.Client | None = None, min_interval: float = 1.1) -> Path:
    args = parse_args(argv)
    keys = credentials()
    if args.from_collector:
        collected = read_collector_csv(args.from_collector)
        site = urlparse(collected[0].get("Product URL") or "").hostname if collected else None
        marketplace, query, notes = args.marketplace or site or "www.amazon.eg", {}, []
        use_api = bool(keys) and not args.no_api
        if not use_api:
            notes.append("exported the collected page data only" + ("" if args.no_api else " (no Creators API credentials set)"))
    else:
        marketplace, query, notes = query_from_url(args.url) if args.url else ("www.amazon.eg", {}, [])
        marketplace = args.marketplace or marketplace
        for key, value in (("keywords", args.keywords), ("browseNodeId", args.browse_node),
                           ("sortBy", args.sort), ("searchIndex", args.search_index)):
            if value:
                query[key] = value
        if not (query.get("keywords") or query.get("browseNodeId")):
            raise SystemExit("the link has no search words (k=) or category (rh=n:...); add --keywords or --browse-node")
        if not keys:
            raise SystemExit(f"set {', '.join(CREDENTIAL_VARIABLES)} (from Associates Central > Creators API) before running")
        use_api = True

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    out = args.out or Path(__file__).resolve().parents[1] / "Data" / "amazon" / f"{marketplace.removeprefix('www.')}_{stamp}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    raw_path = out.with_suffix(".raw.jsonl")
    for note in notes:
        print(f"note: {note}", file=sys.stderr)

    total = None
    details: dict[str, dict] = {}
    items: list[dict] = []
    if use_api:
        with (http or httpx.Client(timeout=30)) as client, raw_path.open("w", encoding="utf-8") as raw_log:
            api = CreatorsApi(*keys, marketplace, client, token_url=os.environ.get("CREATORS_TOKEN_URL"),
                              min_interval=min_interval, raw_log=raw_log)
            if args.from_collector:
                print(f"fetching details of {len(collected)} products from {marketplace}", file=sys.stderr)
                details = get_items(api, [row["ASIN"] for row in collected])
            else:
                print(f"searching {marketplace} for {json.dumps(query, ensure_ascii=False)}", file=sys.stderr)
                items, total = (search_by_price_slices if args.price_slices else search)(api, query, args.pages)
                if args.variations:
                    print("fetching variations", file=sys.stderr)
                    items = expand_variations(api, items)

    spec_columns: list[str] = []
    if args.from_collector:
        rows = [merge_collected(row, details.get(row["ASIN"])) for row in collected]
        spec_columns = add_spec_columns(rows)
        pages = sorted({row["Page"] for row in collected if isinstance(row.get("Page"), int)})
        run_info = {"Collector CSV": args.from_collector.resolve(), "Marketplace": marketplace,
                    "Search pages collected": ", ".join(map(str, pages)) or None, "Products exported": len(rows),
                    "Products with page details": sum(1 for row in rows if row.get("Details collected at")),
                    "Spec columns": len(spec_columns), "Enriched through the API": len(details) if use_api else "no"}
    else:
        rows = list({row["ASIN"]: row for row in map(product_row, items)}.values())
        run_info = {"Source link": args.url, "Marketplace": marketplace, "API query": json.dumps(query, ensure_ascii=False),
                    "Price slices": args.price_slices, "Results Amazon reports": total, "Products exported": len(rows),
                    "Variations expanded": args.variations}
    run_info.update({"Notes": "\n".join(notes) or None, "Exported at": datetime.now().isoformat(timespec="seconds"),
                     "Raw API responses": raw_path.name if use_api else None})
    write_workbook(out, rows, run_info, spec_columns)
    print(f"wrote {len(rows)} products to {out}", file=sys.stderr)
    return out


if __name__ == "__main__":
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")  # Arabic search words on a cp1252 console
    try:
        main()
    except (ApiError, httpx.HTTPError) as error:
        raise SystemExit(f"error: {error}")
